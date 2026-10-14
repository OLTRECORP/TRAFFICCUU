from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

from app.nombres import Indice, palabras, parecido
from app.remisiones_excel import importar, lunes, parse_libro
from app.semanal import emparejar, es_del_cliente


# ---- nombres ---------------------------------------------------------------------------------------------------
def test_variantes_de_captura_de_lugares():
    idx = Indice()
    for i, n in enumerate(["CRIBA TAVO", "ALAMAR", "HOLCIM SIMON BOLIVAR", "HOLCIM EL REFUGIO", "RUMOROSA", "VENADOS", "YARDA TECATE", "CAÑON DE SAENZ"]):
        idx.agregar(n, str(i))
    assert idx.buscar("VALLE CRIVA TAVO") == ("0", "parecido")       # una letra distinta + palabra extra
    assert idx.buscar("YARDA ALAMAR") == ("1", "parecido")
    assert idx.buscar("SIMON BOLIVAR") == ("2", "parecido")
    assert idx.buscar("HOLCIM REFUGIO") == ("3", "igual")            # sin artículos
    assert idx.buscar("LA RUMOROSA") == ("4", "igual")
    assert idx.buscar("LOS VENADOS") == ("5", "igual")
    assert idx.buscar("CAÑON DEL SAENZ") == ("7", "igual")
    assert idx.buscar("EL POZO") == (None, "nuevo")
    assert idx.buscar("HOLCIM") == (None, "ambiguo")                 # parcial en dos lugares: no se adivina


def test_choferes_parcial_requiere_dos_palabras():
    idx = Indice(minimo_parcial=2)
    idx.agregar("CARLOS EDUARDO LOPEZ VILLASEÑOR", "c1")
    idx.agregar("ALEJANDRO VARGAS RUIZ", "c2")
    assert idx.buscar("CARLOS LOPEZ") == ("c1", "parecido")
    assert idx.buscar("ALEJANDRO VARGAZ") == ("c2", "parecido")
    assert idx.buscar("CARLOS") == (None, "nuevo")
    assert not parecido(palabras("COX"), palabras("COXA"), 1)        # palabras cortas: sin tolerancia


# ---- Excel -----------------------------------------------------------------------------------------------------
def libro() -> bytes:
    wb = Workbook()
    p = wb.active
    p.title = "PRECIOS"
    p.append(["ORIGEN", "DESTINO", "PRECIO", "UNIDAD DE MEDIDA"])
    p.append(["COX", "ALAMAR", 160, "M3"])
    p.append(["RUMOROSA", "LOS VENADOS", 395, "M3"])
    v = wb.create_sheet("VIAJES")
    for _ in range(3):
        v.append([None])
    v.append(["SEMANA", "FECHA", "ORIGEN", "DESTINO", "REMISION", "MATERIAL", "NOMBRE OPERADOR", "PLACAS CAMION", "PLACA GONDOLA",
              "PLACA GONDOLA 2", "CANTIDAD", "UNIDAD", "PRECIO UNITARIO", "SUBTOTAL"])
    v.append([None, datetime(2025, 7, 7), "COX", "YARDA ALAMAR", 1, "ARENA", "JUAN PEREZ", "80A1PM", "89UP8A ", None, 32, "m³", 160, 5120])
    v.append([None, datetime(2025, 7, 8), "LA RUMOROSA", "VENADOS", 2, "SILICA", "JUAN PEREZ", "80a1pm", "89UP8A", None, 32, "M3", 395, 12640])
    v.append([None, datetime(2025, 7, 9), None, "VENADOS", "N/A", "SILICA", "JUAN PEREZ", "80A1PM", "89UP8A", None, 32, "M3", 395, 12640])
    v.append([None, datetime(2025, 7, 9), "COX", "ALAMAR", 1, "ARENA", "JUAN PEREZ", "80A1PM", "89UP8A", None, 32, "M3", 170, 5440])
    v.append([None, "N/A", None, None, None, None, None, None, None, None, None, None, None, 0])
    v.append([None, datetime(2025, 1, 10), "COX", "ALAMAR", 3, "ARENA", "JUAN PEREZ", "80A1PM", "89UP8A", None, 32, "M3", 160, 5120])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_parse_libro_encabezado_desplazado_y_omitidas():
    r = parse_libro(libro())
    assert len(r["precios"]) == 2 and len(r["filas"]) == 5
    assert r["omitidas"] == [{"fila": 9, "motivo": "sin fecha válida"}]
    f = r["filas"][0]
    assert (f.unidad, f.tracto, f.gondola, f.remision) == ("M3", "80A1PM", "89UP8A", "1")
    assert r["filas"][2].remision is None and r["filas"][2].origen is None


class _Q:
    def __init__(self, db, t):
        self.db, self.t, self.filtros, self.op, self.p = db, t, [], "select", None
    def select(self, *_): return self
    def order(self, *_): return self
    def limit(self, *_): return self
    def range(self, a, b): self.rng = (a, b); return self
    def eq(self, k, v): self.filtros.append((k, v)); return self
    def insert(self, p): self.op, self.p = "insert", p; return self
    def update(self, p): self.op, self.p = "update", p; return self
    def upsert(self, p, **_): self.op, self.p = "upsert", p; return self
    def execute(self):
        filas = self.db.t.setdefault(self.t, [])
        if self.op in ("insert", "upsert"):
            nuevos = self.p if isinstance(self.p, list) else [self.p]
            for n in nuevos:
                filas.append({**n, "id": n.get("id") or f"{self.t}{len(filas)}"})
            return type("R", (), {"data": filas[-len(nuevos):]})()
        sel = [f for f in filas if all(f.get(k) == v for k, v in self.filtros if k != "empresa_id")]
        if self.op == "update":
            for f in sel:
                f.update(self.p)
        if getattr(self, "rng", None) and self.rng[0] > 0:
            sel = []
        return type("R", (), {"data": sel})()


class FakeDB:
    def __init__(self, **t): self.t = t
    def table(self, n): return _Q(self, n)


def test_importar_vista_previa_y_escritura():
    db = FakeDB(lugares=[{"id": "L1", "nombre": "ALAMAR", "tipo": "destino", "alias": [], "cliente": None}],
                choferes=[{"id": "C1", "nombre": "JUAN PEREZ LOPEZ", "rfc": "PELJ800101AAA"}],
                tractocamiones=[{"id": "T1", "placa": "80A1PM"}], gondolas=[], unidad_alias=[], viajes=[], rutas=[], clientes=[])
    parsed = parse_libro(libro())
    prev = importar(db, "e1", parsed, None, confirmar=False)
    assert prev["nuevos"] == 5 and prev["cliente"]["nombre"] == "PETREOS LA PURISIMA"
    assert not db.t["viajes"] and not db.t["clientes"]                            # la vista previa no escribe
    assert {(l["excel"], l["lugar"], l["accion"]) for l in prev["lugares"]} >= {
        ("YARDA ALAMAR", "ALAMAR", "parecido"), ("LA RUMOROSA", "RUMOROSA", "variante")}
    assert prev["avisos"]["origen_deducido"][0]["origen"] == "RUMOROSA"
    assert prev["avisos"]["precio_distinto_tarifa"][0]["precio"] == 170.0
    assert prev["avisos"]["fecha_fuera_de_secuencia"][0]["fecha"] == "2025-01-10"
    assert prev["choferes"][0]["accion"] == "parecido"                           # JUAN PEREZ -> JUAN PEREZ LOPEZ

    r = importar(db, "e1", parse_libro(libro()), None, confirmar=True)
    assert r["nuevos"] == 5 and len(db.t["viajes"]) == 5
    v = db.t["viajes"][0]
    assert (v["origen_registro"], v["tracto_id"], v["chofer_id"], v["lugar_destino_id"]) == ("excel", "T1", "C1", "L1")
    c = db.t["clientes"][0]
    assert (c["comision_pct"], c["km_por_mapa"], c["factura_semanal"]) == (11, True, True)
    rumorosa = next(t for t in db.t["tarifas"] if t.get("comision_pct") == 10)
    assert rumorosa["precio"] == 395.0
    assert importar(db, "e1", parse_libro(libro()), c["id"], confirmar=True)["nuevos"] == 0   # reimportar no duplica


# ---- facturas semanales -------------------------------------------------------------------------------------------
def test_emparejar_por_monto_y_fecha():
    semanas = [
        {"semana": "2025-07-07", "fin": "2025-07-13", "subtotal": 1000.0, "facturas": {}},
        {"semana": "2025-07-14", "fin": "2025-07-20", "subtotal": 1000.0, "facturas": {}},
        {"semana": "2025-07-21", "fin": "2025-07-27", "subtotal": 500.0, "facturas": {"F9": 3}},
    ]
    facturas = [{"id": "F1", "fecha": "2025-07-22", "subtotal": 1000.4},   # cuadra con dos: gana la más cercana (14-20)
                {"id": "F2", "fecha": "2025-07-23", "subtotal": 1000.0},   # toma la que queda (7-13)
                {"id": "F3", "fecha": "2025-07-23", "subtotal": 777.0},    # no cuadra con ninguna
                {"id": "F9", "fecha": "2025-07-28", "subtotal": 500.0}]    # ya ligada
    pares = {f["id"]: s["semana"] for f, s in emparejar(semanas, facturas)}
    assert pares == {"F1": "2025-07-14", "F2": "2025-07-07"}


def test_factura_anterior_a_la_semana_no_se_liga():
    semanas = [{"semana": "2025-07-14", "fin": "2025-07-20", "subtotal": 100.0, "facturas": {}}]
    assert emparejar(semanas, [{"id": "F", "fecha": "2025-07-10", "subtotal": 100.0}]) == []


def test_cliente_por_rfc_o_nombre():
    c = {"nombre": "PETREOS LA PURISIMA", "rfc": "PPU010101AB1", "alias": ["PURISIMA"]}
    assert es_del_cliente(c, "ppu010101ab1", None)
    assert es_del_cliente({**c, "rfc": None}, None, "Pétreos La Purísima, S.A. de C.V.")
    assert not es_del_cliente({**c, "rfc": None}, "XAXX010101000", "TRITURADOS ROCA")
    assert lunes(datetime(2025, 7, 13).date()).isoformat() == "2025-07-07"


def test_viajes_casi_iguales_con_distinto_folio_se_importan_ambos_sin_aviso():
    wb = Workbook()
    wb.active.title = "PRECIOS"
    wb.active.append(["ORIGEN", "DESTINO", "PRECIO", "UNIDAD"])
    wb.active.append(["COX", "ALAMAR", 160, "M3"])
    v = wb.create_sheet("VIAJES")
    v.append(["FECHA", "ORIGEN", "DESTINO", "REMISION", "MATERIAL", "NOMBRE OPERADOR", "PLACAS CAMION", "PLACA GONDOLA", "CANTIDAD", "UNIDAD", "PRECIO UNITARIO", "SUBTOTAL"])
    for folio in (100, 101):   # mismo día, ruta, chofer, tracto y precio acordado distinto a la tarifa
        v.append([datetime(2025, 7, 7), "COX", "ALAMAR", folio, "ARENA", "JUAN PEREZ", "80A1PM", "89UP8A", 32, "M3", 170, 5440])
    out = BytesIO(); wb.save(out)
    db = FakeDB(lugares=[], choferes=[], tractocamiones=[], gondolas=[], unidad_alias=[], viajes=[], rutas=[], clientes=[])
    r = importar(db, "e1", parse_libro(out.getvalue()), None, confirmar=True)
    assert r["nuevos"] == 2 and len(db.t["viajes"]) == 2
    assert "remision_repetida" not in r["avisos"]
    assert {v["precio_unitario"] for v in db.t["viajes"]} == {170.0}             # se respeta el precio acordado del viaje
    assert r["avisos"]["precio_distinto_tarifa"][0]["viajes"] == 2               # solo informativo
