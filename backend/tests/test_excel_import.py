from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.excel_import import ExcelError, guardar, parse_workbook

U1 = "ED72CE46-0DA3-4E3C-B998-90D63C327EB0"
U2 = "4AB096DF-B3FD-4953-8130-A6AE4C0F89AE"


def make_xlsx(facturas=None, ordenes=None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    if facturas is not None:
        ws = wb.create_sheet("LISTA DE FACTURAS")
        ws.append(["ULTIMOS 5", "C.PAG.", "Folio", "UUID", "Tipo de comprobante", "Estado", "Fecha de emision", "Total",
                   "COMPLEMENTO", "ORDEN REFERIDA", "ORDEN DE COMPRA", "FECHA DE CREACION", "FECHA DE PAGO"])
        for r in facturas:
            ws.append(r)
    if ordenes is not None:
        ws = wb.create_sheet("RELACION DE ORDENES")
        ws.append(["N° DE ORDEN", "MONTO", "FECHA DE CREACION", "FECHA DE PAGO", "SUMA FACTURAS EN ORDEN", "COMP PAGO"])
        for r in ordenes:
            ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_facturas_y_ordenes():
    data = make_xlsx(
        facturas=[
            ["27EB0", 1, 3319, U1, "factura", "Activado", "08/01/2026", 1515.63, "COMPLEMENTO LISTO", "4287CH", "4287CH",
             datetime(2026, 1, 15), datetime(2026, 2, 13)],
            ["F89AE", 0, 3320, U2, "factura", "activado", "23/03/2026", 10, "SIN COMPLEMENTO", 0, None, "NO ASIGNADO", " SIN ORDEN"],
            ["x", 0, 1, "no-es-uuid", "factura", "activado", None, 1, None, None, None, None, None],
            ["dup", 1, 3319, U1.lower(), "factura", "pagado", "08/01/2026", 2000, None, None, "4287CH", None, None],
        ],
        ordenes=[
            ["4287CH", 117308.24, datetime(2026, 1, 15), datetime(2026, 2, 13), 1, 117308.23],
            [None, None, None, "-", 0, 0],
            ["TOTAL EN ORDEN", 6115411.16, None, "NO. DE FACTURAS", 1, 1],
        ],
    )
    p = parse_workbook(data)
    assert len(p["facturas"]) == 3  # folio repetido se fusiona; UUID dañado se conserva sin UUID
    f1 = next(f for f in p["facturas"] if f["folio"] == "3319")
    assert f1["estado"] == "pagado" and f1["total"] == "2000.00" and f1["uuid"] == U1.lower()
    assert next(f for f in p["facturas"] if f["folio"] == "1")["uuid"] is None
    assert p["uuid_danados"] == 1
    f2 = next(f for f in p["facturas"] if f["folio"] == "3320")
    assert f2["orden_referida"] is None and f2["orden_compra"] is None
    assert f2["fecha_emision"] == "2026-03-23" and f2["fecha_creacion"] is None and f2["fecha_pago"] is None
    assert p["ordenes"] == [{"numero": "4287CH", "monto": "117308.24", "fecha_creacion": "2026-01-15",
                             "fecha_pago": "2026-02-13", "comp_pago": "117308.23"}]
    assert p["omitidas"] == 1  # renglón de totales


def test_hojas_faltantes():
    wb = Workbook()
    wb.active.title = "OTRA"
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(ExcelError):
        parse_workbook(buf.getvalue())
    with pytest.raises(ExcelError):
        parse_workbook(b"no es un excel")


class FakeTable:
    def __init__(self, store, name):
        self.store, self.name, self._in = store, name, None

    def select(self, *_): return self
    def eq(self, *_): return self
    def in_(self, key, keys): self._in = (key, keys); return self
    def order(self, *_): return self
    def range(self, a, b): self._rango = (a, b); return self

    def upsert(self, rows, on_conflict):
        self.key = on_conflict.split(",")[1]
        for r in rows:
            t = self.store.setdefault(self.name, {})
            t[r[self.key]] = {"id": f"{self.name}-{r[self.key]}", **t.get(r[self.key], {}), **r}
        return self

    def execute(self):
        class R: pass
        r = R()
        if self._in:
            key, keys = self._in
            r.data = [{key: k} for k in keys if k in self.store.get(self.name, {})]
        elif getattr(self, "_rango", None):
            a, b = self._rango
            r.data = list(self.store.get(self.name, {}).values())[a:b + 1]
        else:
            r.data = []
        return r


class FakeDb:
    def __init__(self): self.store = {}
    def table(self, name): return FakeTable(self.store, name)


def test_guardar_es_idempotente():
    parsed = parse_workbook(make_xlsx(
        facturas=[["a", 1, 1, U1, "factura", "activado", "08/01/2026", 5, None, None, "4287CH", None, None]],
        ordenes=[["4287CH", 5, None, None, 5, 5]],
    ))
    db = FakeDb()
    r1 = guardar(db, "emp", parsed)
    assert r1["facturas"] == {"total": 1, "nuevas": 1, "actualizadas": 0}
    r2 = guardar(db, "emp", parsed)
    assert r2["facturas"] == {"total": 1, "nuevas": 0, "actualizadas": 1}
    assert r2["ordenes_compra"]["actualizadas"] == 1
    assert len(db.store["facturas"]) == 1


def test_excel_posterior_actualiza_la_factura_creada_desde_el_xml_sin_duplicar():
    # El XML ya dio de alta la factura (folio "123", con UUID); el Excel la trae con el UUID dañado y folio "0123".
    db = FakeDb()
    db.store["facturas"] = {"123": {"id": "f1", "folio": "123", "uuid": U1, "total": "5000.00",
                                       "tipo_comprobante": "Ingreso", "fecha_emision": "2026-01-08"}}
    parsed = parse_workbook(make_xlsx(
        facturas=[["a", 1, "0123", "UUID DANADO", "factura", "pagado", None, 5000, None, None, "4287CH", None, None]],
        ordenes=[],
    ))
    fila = parsed["facturas"][0]
    assert fila["folio"] == "0123" and fila["uuid"] is None
    r = guardar(db, "emp", parsed)
    assert r["facturas"] == {"total": 1, "nuevas": 0, "actualizadas": 1}
    assert list(db.store["facturas"]) == ["123"]                 # sin duplicar
    f = db.store["facturas"]["123"]
    assert f["uuid"] == U1 and f["estado"] == "pagado" and f["fecha_emision"] == "2026-01-08"
