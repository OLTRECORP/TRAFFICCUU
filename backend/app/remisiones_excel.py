"""Importación de viajes por remisión desde Excel (formato de Pétreos La Purísima).

Hojas (por nombre, sin importar mayúsculas):
  - PRECIOS: ORIGEN, DESTINO, PRECIO, UNIDAD DE MEDIDA                  -> tarifas del cliente por ruta
  - VIAJES:  FECHA, ORIGEN, DESTINO, REMISION, MATERIAL, NOMBRE OPERADOR, PLACAS CAMION, PLACA GONDOLA,
             CANTIDAD, UNIDAD, PRECIO UNITARIO, SUBTOTAL (el encabezado puede no estar en la primera fila)

Cada renglón de VIAJES es un viaje (sin XML). Orígenes y destinos se ligan al catálogo de lugares tolerando
variantes de captura («LA RUMOROSA» = «RUMOROSA», «VALLE CRIVA TAVO» = «CRIBA TAVO»); los nuevos se crean sin
ubicación para marcarla en el mapa. Choferes y unidades se buscan por nombre y placa y se crean si faltan.
Reimportar el mismo archivo no duplica: cada renglón lleva una clave estable (`import_key`).
`importar(..., confirmar=False)` solo devuelve la vista previa, sin escribir nada."""
import hashlib
import logging
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

from .nombres import Indice, normalizar, palabras

log = logging.getLogger("remisiones")

LOTE = 200
DIAS_FUERA_DE_SECUENCIA = 120   # una fecha tan anterior a las ya vistas suele ser un año mal capturado

# Cliente que se crea si se importa sin elegir uno, con lo que pidió el usuario para La Purísima:
# comisión 11 %, km por mapa, factura semanal (lunes a domingo) y 10 % en la ruta Rumorosa -> Venados (arena sílica).
PURISIMA = {"nombre": "PETREOS LA PURISIMA", "alias": ["LA PURISIMA", "PURISIMA"], "comision_pct": 11,
            "km_por_mapa": True, "factura_semanal": True}
PURISIMA_EXCEPCIONES = [({"rumorosa"}, {"venados"}, 10)]

COLUMNAS = {
    "fecha": ("fecha",), "origen": ("origen",), "destino": ("destino",), "remision": ("remision", "remisión"),
    "material": ("material",), "operador": ("nombre operador", "operador", "chofer"),
    "tracto": ("placas camion", "placa camion", "placas tracto", "placa tracto"),
    "gondola": ("placa gondola", "placas gondola"), "cantidad": ("cantidad",), "unidad": ("unidad",),
    "precio": ("precio unitario", "precio"), "subtotal": ("subtotal", "importe"),
}


class RemisionesError(ValueError):
    pass


@dataclass
class Fila:
    fila: int
    fecha: date
    origen: str | None
    destino: str
    remision: str | None
    material: str | None
    operador: str | None
    tracto: str | None
    gondola: str | None
    cantidad: Decimal | None
    unidad: str | None
    precio: Decimal | None
    subtotal: Decimal


def _texto(v) -> str | None:
    if v is None:
        return None
    t = re.sub(r"\s+", " ", str(v)).strip()
    return t or None


def _dec(v) -> Decimal | None:
    if v is None or isinstance(v, bool) or v == "":
        return None
    try:
        return Decimal(str(v).replace(",", "").replace("$", "").strip())
    except (InvalidOperation, ValueError):
        return None


def _fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = _texto(v)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime((t or "")[:10], fmt).date()
        except ValueError:
            continue
    return None


def placa(v) -> str | None:
    t = re.sub(r"[^A-Z0-9]", "", str(v or "").upper())
    return t or None


def unidad(v) -> str | None:
    t = _texto(v)
    if not t:
        return None
    return "M3" if normalizar(t.replace("³", "3")) in ("m3", "m 3", "metro cubico", "metros cubicos") else t.upper()


def _hoja(wb, nombre: str):
    return next((ws for ws in wb.worksheets if normalizar(ws.title) == nombre), None)


def _encabezado(ws) -> tuple[int, dict[str, int]]:
    """Fila del encabezado (la primera con FECHA, ORIGEN y DESTINO) y columna de cada campo."""
    for i, fila in enumerate(ws.iter_rows(min_row=1, max_row=40, values_only=True), start=1):
        nombres = [normalizar(_texto(c)) for c in fila]
        if {"fecha", "origen", "destino"} <= set(nombres):
            cols: dict[str, int] = {}
            for campo, variantes in COLUMNAS.items():
                for j, n in enumerate(nombres):
                    if n in {normalizar(x) for x in variantes} and campo not in cols:
                        cols[campo] = j
            return i, cols
    raise RemisionesError("La hoja VIAJES no tiene un encabezado con FECHA, ORIGEN y DESTINO")


def parse_libro(contenido: bytes) -> dict:
    try:
        wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:
        raise RemisionesError(f"No se pudo leer el Excel: {exc}") from exc

    precios = []
    hp = _hoja(wb, "precios")
    if hp is not None:
        for fila in hp.iter_rows(min_row=2, values_only=True):
            o, d = _texto(fila[0] if len(fila) > 0 else None), _texto(fila[1] if len(fila) > 1 else None)
            p = _dec(fila[2] if len(fila) > 2 else None)
            if o and d and p is not None:
                precios.append({"origen": o, "destino": d, "precio": p, "unidad": unidad(fila[3] if len(fila) > 3 else None) or "M3"})

    hv = _hoja(wb, "viajes")
    if hv is None:
        raise RemisionesError("El Excel no tiene la hoja VIAJES")
    inicio, cols = _encabezado(hv)
    faltan = [c for c in ("fecha", "destino", "subtotal") if c not in cols]
    if faltan:
        raise RemisionesError(f"Faltan columnas en VIAJES: {', '.join(faltan)}")

    filas, omitidas = [], []
    for n, r in enumerate(hv.iter_rows(min_row=inicio + 1, values_only=True), start=inicio + 1):
        def v(campo):
            j = cols.get(campo)
            return r[j] if j is not None and j < len(r) else None
        if all(c is None or _texto(c) is None for c in r):
            continue
        f, destino, subtotal = _fecha(v("fecha")), _texto(v("destino")), _dec(v("subtotal"))
        motivo = ("sin fecha válida" if f is None else "sin destino" if not destino
                  else "sin subtotal" if subtotal is None else "subtotal en cero" if subtotal == 0 else None)
        if motivo:
            omitidas.append({"fila": n, "motivo": motivo})
            continue
        rem = _texto(v("remision"))
        filas.append(Fila(
            fila=n, fecha=f, origen=_texto(v("origen")), destino=destino,
            remision=None if rem and normalizar(rem) in ("n a", "na", "sin remision") else rem,
            material=_texto(v("material")), operador=_texto(v("operador")),
            tracto=placa(v("tracto")), gondola=placa(v("gondola")),
            cantidad=_dec(v("cantidad")), unidad=unidad(v("unidad")), precio=_dec(v("precio")), subtotal=subtotal,
        ))
    if not filas:
        raise RemisionesError("La hoja VIAJES no tiene viajes con fecha, destino y subtotal")
    return {"precios": precios, "filas": filas, "omitidas": omitidas}


def lunes(d: date) -> date:
    return d - timedelta(days=d.weekday())


def clave_import(f: Fila, ocurrencia: int) -> str:
    base = "|".join(str(x) for x in (
        f.fecha.isoformat(), normalizar(f.remision), normalizar(f.origen), normalizar(f.destino),
        f.tracto or "", normalizar(f.operador), f.cantidad or "", f.precio or "", f.subtotal))
    return "rem:" + hashlib.sha1(base.encode()).hexdigest()[:24] + f"#{ocurrencia}"


def _leer(db, tabla: str, columnas: str, empresa_id: str) -> list[dict]:
    filas, desde = [], 0
    while True:
        lote = db.table(tabla).select(columnas).eq("empresa_id", empresa_id).order("id").range(desde, desde + 999).execute().data or []
        filas += lote
        if len(lote) < 1000:
            return filas
        desde += 1000


class _Plan:
    """Resuelve nombres y placas contra los catálogos; en vista previa los nuevos llevan un id provisional."""

    def __init__(self, db, empresa_id: str, cliente: dict, confirmar: bool):
        self.db, self.eid, self.cliente, self.confirmar = db, empresa_id, cliente, confirmar
        self.lugares = {l["id"]: l for l in _leer(db, "lugares", "id, nombre, tipo, alias, cliente", empresa_id)}
        self.idx_global, self.idx_cliente = Indice(), Indice()
        for l in self.lugares.values():
            for n in [l["nombre"], *(l.get("alias") or [])]:
                self.idx_global.agregar(n, l["id"])
                if normalizar(l.get("cliente")) == normalizar(cliente["nombre"]):
                    self.idx_cliente.agregar(n, l["id"])
        self.mapa_lugares: dict[str, dict] = {}
        self.choferes = {c["id"]: c for c in _leer(db, "choferes", "id, nombre, rfc", empresa_id)}
        self.idx_choferes = Indice(minimo_parcial=2)
        for c in self.choferes.values():
            self.idx_choferes.agregar(c["nombre"], c["id"])
        self.mapa_choferes: dict[str, dict] = {}
        self.unidades = {"tracto": {}, "gondola": {}}
        for tipo, tabla in (("tracto", "tractocamiones"), ("gondola", "gondolas")):
            for u in _leer(db, tabla, "id, placa", empresa_id):
                if placa(u["placa"]):
                    self.unidades[tipo][placa(u["placa"])] = {"id": u["id"], "accion": "existente"}
        for a in _leer(db, "unidad_alias", "id, tipo, placa, unidad_id", empresa_id):
            if placa(a["placa"]):
                self.unidades[a["tipo"]].setdefault(placa(a["placa"]), {"id": a["unidad_id"], "accion": "existente"})
        self.nuevos = Counter()

    # -- lugares -----------------------------------------------------------------------------------------------
    def lugar(self, nombre: str, tipo: str) -> str:
        clave = normalizar(nombre)
        m = self.mapa_lugares.get(clave)
        if m is None:
            ident, como = self.idx_global.buscar(nombre)
            if ident is None or como != "igual":
                ident, como = self.idx_cliente.buscar(nombre)
            if ident is None:
                ident, como = self._nuevo_lugar(nombre, tipo), "nuevo"
            elif normalizar(nombre) != normalizar(self.lugares[ident]["nombre"]):
                self._alias(ident, nombre)
                como = "parecido" if como == "parecido" else "variante"   # «LA RUMOROSA» = «RUMOROSA»
            if como in ("igual", "variante"):   # un lugar ya existente que usa este cliente: sus variantes se comparan contra él
                self.idx_cliente.agregar(self.lugares[ident]["nombre"], ident)
            m = self.mapa_lugares[clave] = {"excel": nombre, "id": ident, "lugar": self.lugares[ident]["nombre"],
                                            "accion": "existente" if como == "igual" else como, "usos": set()}
        m["usos"].add(tipo)
        l = self.lugares[m["id"]]
        if l["tipo"] not in (tipo, "ambos"):
            l["tipo"] = "ambos"
            if self.confirmar and not m["id"].startswith("nuevo:"):
                self.db.table("lugares").update({"tipo": "ambos"}).eq("id", m["id"]).execute()
        return m["id"]

    def _nuevo_lugar(self, nombre: str, tipo: str) -> str:
        fila = {"empresa_id": self.eid, "nombre": nombre.upper(), "tipo": tipo, "cliente": self.cliente["nombre"],
                "auto_ubicar": False, "referencia": "Creado desde el Excel de remisiones: marca su ubicación en el mapa"}
        if self.confirmar:
            # si el lugar ya existe con otro cliente (mismo nombre), se reutiliza en lugar de chocar con el único
            previo = self.db.table("lugares").select("id, nombre, tipo, alias, cliente").eq("empresa_id", self.eid).eq("nombre", fila["nombre"]).limit(1).execute().data
            l = previo[0] if previo else self.db.table("lugares").insert(fila).execute().data[0]
        else:
            l = {**fila, "id": f"nuevo:{normalizar(nombre)}", "alias": []}
        l.setdefault("alias", [])
        self.lugares[l["id"]] = l
        self.idx_cliente.agregar(l["nombre"], l["id"])
        self.idx_global.agregar(l["nombre"], l["id"])
        self.nuevos["lugares"] += 1
        return l["id"]

    def _alias(self, ident: str, nombre: str) -> None:
        l = self.lugares[ident]
        if normalizar(nombre) == normalizar(l["nombre"]) or any(normalizar(a) == normalizar(nombre) for a in l.get("alias") or []):
            return
        l["alias"] = [*(l.get("alias") or []), nombre.upper()]
        self.idx_cliente.agregar(nombre, ident)
        if self.confirmar and not ident.startswith("nuevo:"):
            self.db.table("lugares").update({"alias": l["alias"]}).eq("id", ident).execute()

    # -- choferes y unidades -------------------------------------------------------------------------------------
    def preparar_choferes(self, nombres: set[str]) -> None:
        # los nombres más completos primero: «CARLOS LOPEZ» se liga a «CARLOS EDUARDO LOPEZ VILLASEÑOR» si ya existe
        for nombre in sorted(nombres, key=lambda n: (-len(palabras(n)), n)):
            ident, como = self.idx_choferes.buscar(nombre)
            if ident is None:
                fila = {"empresa_id": self.eid, "nombre": nombre.upper(),
                        "rfc": "SINRFC" + hashlib.sha1(normalizar(nombre).encode()).hexdigest()[:7].upper(),
                        "notas": f"Creado desde el Excel de remisiones de {self.cliente['nombre']}: falta capturar su RFC"}
                if self.confirmar:
                    previo = self.db.table("choferes").select("id, nombre, rfc").eq("empresa_id", self.eid).eq("rfc", fila["rfc"]).limit(1).execute().data
                    c = previo[0] if previo else self.db.table("choferes").insert(fila).execute().data[0]
                else:
                    c = {**fila, "id": f"nuevo:{normalizar(nombre)}"}
                self.choferes[c["id"]] = c
                self.idx_choferes.agregar(c["nombre"], c["id"])
                ident, como = c["id"], "nuevo"
                self.nuevos["choferes"] += 1
            self.mapa_choferes[normalizar(nombre)] = {"excel": nombre, "id": ident, "chofer": self.choferes[ident]["nombre"],
                                                      "accion": "existente" if como == "igual" else como}

    def chofer(self, nombre: str | None) -> str | None:
        return self.mapa_choferes[normalizar(nombre)]["id"] if nombre else None

    def unidad(self, tipo: str, p: str | None) -> str | None:
        if not p:
            return None
        u = self.unidades[tipo].get(p)
        if u is None:
            tabla = "tractocamiones" if tipo == "tracto" else "gondolas"
            if self.confirmar:
                u = {"id": self.db.table(tabla).insert({"empresa_id": self.eid, "placa": p}).execute().data[0]["id"], "accion": "nuevo"}
            else:
                u = {"id": f"nuevo:{p}", "accion": "nuevo"}
            self.unidades[tipo][p] = u
            self.nuevos["tractos" if tipo == "tracto" else "gondolas"] += 1
        return u["id"]


def asegurar_cliente(db, empresa_id: str, cliente_id: str | None, confirmar: bool) -> tuple[dict, bool]:
    """(cliente, creado). Sin cliente elegido se usa (o crea) Pétreos La Purísima."""
    if cliente_id:
        c = db.table("clientes").select("*").eq("empresa_id", empresa_id).eq("id", cliente_id).limit(1).execute().data
        if not c:
            raise RemisionesError("Cliente no encontrado")
        return c[0], False
    c = db.table("clientes").select("*").eq("empresa_id", empresa_id).eq("nombre", PURISIMA["nombre"]).limit(1).execute().data
    if c:
        return c[0], False
    fila = {"empresa_id": empresa_id, **PURISIMA}
    if not confirmar:
        return {**fila, "id": "nuevo:cliente", "rfc": None}, True
    return db.table("clientes").insert(fila).execute().data[0], True


def importar(db, empresa_id: str, parsed: dict, cliente_id: str | None = None, confirmar: bool = False) -> dict:
    cliente, cliente_creado = asegurar_cliente(db, empresa_id, cliente_id, confirmar)
    plan = _Plan(db, empresa_id, cliente, confirmar)
    filas: list[Fila] = parsed["filas"]
    omitidas = list(parsed["omitidas"])
    avisos: dict[str, list] = defaultdict(list)

    # 1. Tarifas (los nombres de la hoja PRECIOS se resuelven primero: son la referencia de los del viaje)
    tarifas: dict[tuple[str, str], dict] = {}
    for p in parsed["precios"]:
        o, d = plan.lugar(p["origen"], "origen"), plan.lugar(p["destino"], "destino")
        tarifas[(o, d)] = {"precio": p["precio"], "unidad": p["unidad"], "origen": p["origen"], "destino": p["destino"]}

    # 2. Origen faltante: si una sola tarifa llega a ese destino con ese precio, de ahí salió el viaje
    validas: list[tuple[Fila, str, str]] = []
    for f in filas:
        d = plan.lugar(f.destino, "destino")
        if f.origen:
            o = plan.lugar(f.origen, "origen")
        else:
            posibles = [k for k, t in tarifas.items() if k[1] == d and f.precio is not None and t["precio"] == f.precio]
            if len(posibles) != 1:
                omitidas.append({"fila": f.fila, "motivo": "sin origen (no se pudo deducir de las tarifas)"})
                continue
            o = posibles[0][0]
            f.origen = plan.lugares[o]["nombre"]
            avisos["origen_deducido"].append({"fila": f.fila, "origen": f.origen, "destino": f.destino})
        validas.append((f, o, d))

    plan.preparar_choferes({f.operador for f, _, _ in validas if f.operador})

    # 3. Revisiones de captura
    vistos_max: date | None = None
    rem = Counter(normalizar(f.remision) for f, _, _ in validas if f.remision)
    precio_dif: dict[tuple, int] = Counter()
    sin_tarifa: dict[tuple, int] = Counter()
    for f, o, d in validas:
        if vistos_max and f.fecha < vistos_max - timedelta(days=DIAS_FUERA_DE_SECUENCIA):
            avisos["fecha_fuera_de_secuencia"].append({"fila": f.fila, "fecha": f.fecha.isoformat(), "remision": f.remision})
        vistos_max = max(vistos_max or f.fecha, f.fecha)
        if f.cantidad is not None and f.precio is not None and abs(f.cantidad * f.precio - f.subtotal) > Decimal("0.5"):
            avisos["subtotal_distinto"].append({"fila": f.fila, "cantidad": float(f.cantidad), "precio": float(f.precio), "subtotal": float(f.subtotal)})
        if f.remision and rem[normalizar(f.remision)] > 1:
            avisos["remision_repetida"].append({"fila": f.fila, "remision": f.remision, "fecha": f.fecha.isoformat()})
        t = tarifas.get((o, d))
        ruta = (plan.lugares[o]["nombre"], plan.lugares[d]["nombre"])
        if t is None:
            sin_tarifa[(*ruta, float(f.precio) if f.precio is not None else None)] += 1
        elif f.precio is not None and t["precio"] != f.precio:
            precio_dif[(*ruta, float(f.precio), float(t["precio"]))] += 1
    avisos["precio_distinto_tarifa"] = [{"origen": k[0], "destino": k[1], "precio": k[2], "tarifa": k[3], "viajes": n} for k, n in precio_dif.most_common()]
    avisos["sin_tarifa"] = [{"origen": k[0], "destino": k[1], "precio": k[2], "viajes": n} for k, n in sin_tarifa.most_common()]

    # 4. Viajes
    existentes: set[str] = set()
    if not str(cliente["id"]).startswith("nuevo:"):
        existentes = {v["import_key"] for v in _leer(db, "viajes", "id, import_key", empresa_id) if v.get("import_key")}
    rutas = {(r["origen_id"], r["destino_id"]): float(r["distancia_km"]) for r in _leer(db, "rutas", "id, origen_id, destino_id, distancia_km", empresa_id)}
    ocurrencias: Counter = Counter()
    nuevos: list[dict] = []
    for f, o, d in validas:
        base = clave_import(f, 0)[:-2]
        ocurrencias[base] += 1
        clave = f"{base}#{ocurrencias[base]}"
        if clave in existentes:
            continue
        nuevos.append({
            "empresa_id": empresa_id, "cliente_id": cliente["id"], "origen_registro": "excel", "import_key": clave,
            "fecha": f"{f.fecha.isoformat()}T00:00:00+00:00", "fecha_salida": f"{f.fecha.isoformat()}T00:00:00+00:00",
            "remision": f.remision, "material": f.material, "mercancia": f.material,
            "cantidad": float(f.cantidad) if f.cantidad is not None else None, "unidad": f.unidad,
            "precio_unitario": float(f.precio) if f.precio is not None else None,
            "subtotal": float(f.subtotal), "total": float(f.subtotal),
            "lugar_origen_id": o, "lugar_destino_id": d, "distancia_mapa_km": rutas.get((o, d)),
            "chofer_id": plan.chofer(f.operador), "tracto_id": plan.unidad("tracto", f.tracto),
            "gondola_id": plan.unidad("gondola", f.gondola), "rfc_receptor": cliente.get("rfc"),
            "distancia_xml": 0,
        })

    semanas = Counter(lunes(f.fecha) for f, _, _ in validas)
    resultado = {
        "confirmado": confirmar,
        "cliente": {"id": cliente["id"], "nombre": cliente["nombre"], "creado": cliente_creado},
        "filas": len(filas) + len(parsed["omitidas"]), "validos": len(validas), "nuevos": len(nuevos),
        "ya_importados": len(validas) - len(nuevos), "omitidas": omitidas,
        "semanas": len(semanas), "desde": min(f.fecha for f, _, _ in validas).isoformat() if validas else None,
        "hasta": max(f.fecha for f, _, _ in validas).isoformat() if validas else None,
        "subtotal_nuevos": round(sum(v["subtotal"] for v in nuevos), 2),
        "tarifas": len(tarifas),
        "lugares": sorted(({**{k: v for k, v in m.items() if k != "usos"}, "usos": sorted(m["usos"])} for m in plan.mapa_lugares.values()),
                          key=lambda m: (m["accion"] == "existente", m["lugar"])),
        "choferes": sorted(plan.mapa_choferes.values(), key=lambda m: (m["accion"] == "existente", m["chofer"])),
        "unidades": sorted(({"tipo": t, "placa": p, "accion": u["accion"]} for t in plan.unidades for p, u in plan.unidades[t].items() if u["accion"] == "nuevo"),
                           key=lambda u: (u["tipo"], u["placa"])),
        "nuevos_catalogo": dict(plan.nuevos),
        "avisos": {k: v for k, v in avisos.items() if v},
    }
    if not confirmar:
        return resultado

    # 5. Escritura
    for (o, d), t in tarifas.items():
        db.table("tarifas").upsert({"empresa_id": empresa_id, "cliente_id": cliente["id"], "origen_id": o, "destino_id": d,
                                    "precio": float(t["precio"]), "unidad": t["unidad"]}, on_conflict="cliente_id,origen_id,destino_id").execute()
    if cliente_creado:
        for po, pd, pct in PURISIMA_EXCEPCIONES:
            for (o, d) in tarifas:
                if po <= palabras(plan.lugares[o]["nombre"]) and pd <= palabras(plan.lugares[d]["nombre"]):
                    db.table("tarifas").update({"comision_pct": pct}).eq("cliente_id", cliente["id"]).eq("origen_id", o).eq("destino_id", d).execute()
    for i in range(0, len(nuevos), LOTE):
        db.table("viajes").upsert(nuevos[i:i + LOTE], on_conflict="empresa_id,import_key", ignore_duplicates=True).execute()
    log.info("Remisiones %s: %d viajes nuevos de %d", cliente["nombre"], len(nuevos), len(validas))
    return resultado
